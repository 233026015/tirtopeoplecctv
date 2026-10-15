"""
Tes laporan Excel: memastikan analisis PER KAMERA benar & independen.

Jalankan:  python test_report_analysis.py
Butuh LibreOffice (soffice) untuk menghitung rumus; tanpa itu tes dilewati.
Data di sini SINTETIS, hanya untuk pengujian (tidak masuk laporan produksi).

Cara kerja: dataset dengan jawaban yang sudah diketahui -> build_report_workbook ->
dihitung ulang LibreOffice -> hasil sheet ANALISIS / REKAP PER KAMERA / VALIDASI dibandingkan
dengan hitungan pandas yang INDEPENDEN dari rumus Excel.
"""
import os
import re
import shutil
import subprocess
import tempfile
import unittest
import zipfile
from datetime import date, datetime

import pandas as pd
from openpyxl import load_workbook

import report_builder as rb

SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
NAMA = {1: "CCTV Pintu Penurunan", 2: "CCTV Area Gate", 3: "CCTV Kosong"}
CAMS = pd.DataFrame({"id": [1, 2, 3], "name": [NAMA[1], NAMA[2], NAMA[3]], "location": ["L1", "L2", "L3"]})
COLS = ["id", "timestamp", "person_id", "status", "total_in", "total_out", "camera_id", "snapshot_path"]


def make_df(spec, extra=()):
    """spec: {camera_id: {jam: (n_masuk, n_keluar)}} pada tanggal 2026-10-07."""
    rows, i = [], 0
    for cid, per_hour in spec.items():
        for jam, (nin, nout) in per_hour.items():
            for k in range(nin + nout):
                i += 1
                st = "MASUK" if k < nin else "KELUAR"
                rows.append((i, datetime(2026, 10, 7, jam, k % 60, 0), 1000 + i % 7, st, i, i, cid, None))
    for e in extra:
        i += 1
        rows.append((i,) + tuple(e))
    return pd.DataFrame(rows, columns=COLS)


def recalc(path):
    out = tempfile.mkdtemp()
    subprocess.run([SOFFICE, "--headless", "--convert-to", "xlsx", "--outdir", out, path],
                   capture_output=True, timeout=240, check=True)
    return load_workbook(os.path.join(out, os.path.basename(path)), data_only=True)


def truth(df, cams=CAMS):
    """Hitungan independen (pandas) per kamera."""
    d = df.copy()
    d["ts"] = pd.to_datetime(d["timestamp"], errors="coerce")
    res = {}
    for _, c in cams.iterrows():
        sub = d[d["camera_id"] == c["id"]]
        per = sub.dropna(subset=["ts"]).groupby(sub["ts"].dt.hour).size()
        masuk, keluar = int((sub["status"] == "MASUK").sum()), int((sub["status"] == "KELUAR").sum())
        sepi = None
        if per.empty:
            ramai = None
        else:
            mx = per.max()
            ramai = (int(per[per == mx].index.min()), int(mx))     # seri -> jam paling awal
            # jam sepi: event terendah di antara jam yang punya event (>0); seri -> jam paling awal
            mn = per.min()
            sepi = (int(per[per == mn].index.min()), int(mn))
        res[c["name"]] = dict(ramai=ramai, sepi=sepi, masuk=masuk, keluar=keluar, event=len(sub),
                              uniq=int(sub["person_id"].nunique()), per_hour=per.to_dict())
    return res


def hour_label(h):
    return f"{h:02d}:00 – {h:02d}:59"


class ReportAnalysisTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not SOFFICE:
            raise unittest.SkipTest("LibreOffice (soffice) tidak ditemukan; tes rumus dilewati.")
        cls.tmp = tempfile.mkdtemp()

    def build(self, df, name):
        path = os.path.join(self.tmp, name + ".xlsx")
        rb.build_report_workbook(path, df, CAMS, date(2026, 10, 7), date(2026, 10, 7), snapshot_dir=None)
        return path, recalc(path)

    @staticmethod
    def pola(wb):
        """Baca tabel ANALISIS 'C. POLA PER KAMERA' -> {nama_kamera: baris}."""
        ws = wb["ANALISIS"]
        start = next(r for r in range(1, 120) if str(ws.cell(r, 1).value or "").startswith("C. POLA"))
        out = {}
        for r in range(start + 2, start + 12):
            nm = ws.cell(r, 2).value
            if nm in NAMA.values():
                out[nm] = [ws.cell(r, c).value for c in range(1, 12)]
        return out

    def check_against_truth(self, df, wb):
        t, p = truth(df), self.pola(wb)
        ws_cam = wb["REKAP PER KAMERA"]
        rekap = {ws_cam.cell(r, 2).value: [ws_cam.cell(r, c).value for c in range(1, 15)] for r in range(7, 10)}
        for nm, tv in t.items():
            row = p[nm]   # No, Kamera, JamRamai, EvRamai, JamSepi, EvSepi, Masuk, Keluar, Selisih, IDunik, Ev/ID
            if tv["ramai"] is None:
                self.assertEqual(row[2], rb.NA, f"{nm}: tidak boleh ada jam ramai palsu")
                self.assertEqual(row[3], 0)
            else:
                self.assertEqual(row[2], hour_label(tv["ramai"][0]), f"{nm}: jam ramai")
                self.assertEqual(row[3], tv["ramai"][1], f"{nm}: event jam ramai")
            if tv["sepi"] is None:
                self.assertEqual(row[4], rb.NA, f"{nm}: tidak boleh ada jam sepi palsu")
            else:
                self.assertEqual(row[4], hour_label(tv["sepi"][0]), f"{nm}: jam sepi")
                self.assertEqual(row[5], tv["sepi"][1], f"{nm}: event jam sepi")
            self.assertEqual((row[6], row[7], row[8]), (tv["masuk"], tv["keluar"], tv["masuk"] - tv["keluar"]), nm)
            self.assertEqual(row[9], tv["uniq"], f"{nm}: ID tracking unik")
            # konsisten dengan REKAP PER KAMERA (kolom: D masuk, E keluar, F event, J event tertinggi, K jam tersibuk)
            rk = rekap[nm]
            self.assertEqual((rk[3], rk[4], rk[5]), (tv["masuk"], tv["keluar"], tv["event"]), nm)
            self.assertEqual(rk[9], row[3], f"{nm}: Event Tertinggi harus sama di REKAP PER KAMERA & ANALISIS")
            self.assertEqual(rk[10], row[2], f"{nm}: Jam Tersibuk harus sama di REKAP PER KAMERA & ANALISIS")
        # total event per kamera = jumlah event per jam pada matriks (kecuali event tanpa timestamp)
        ws = wb["ANALISIS"]
        for i, nm in enumerate(NAMA.values()):
            col = 2 + i
            per_hour = [ws.cell(8 + h, col).value for h in range(24)]
            self.assertEqual(per_hour, [t[nm]["per_hour"].get(h, 0) for h in range(24)], f"{nm}: matriks per jam")
        return t, p

    # ------------------------------------------------------------------ skenario
    def test_dua_kamera_jam_ramai_berbeda(self):
        df = make_df({1: {8: (5, 1), 9: (30, 15), 10: (20, 10)},
                      2: {8: (10, 2), 9: (12, 3), 10: (50, 35)}})
        _, wb = self.build(df, "berbeda")
        t, p = self.check_against_truth(df, wb)
        self.assertEqual(p[NAMA[1]][2], "09:00 – 09:59")
        self.assertEqual(p[NAMA[2]][2], "10:00 – 10:59")
        self.assertEqual(p[NAMA[1]][3], 45)
        self.assertEqual(p[NAMA[2]][3], 85)

    def test_jam_ramai_sama_tetap_benar(self):
        df = make_df({1: {8: (3, 1), 9: (40, 10), 10: (5, 5)},
                      2: {8: (7, 1), 9: (20, 6), 10: (2, 2)}})
        _, wb = self.build(df, "sama")
        t, p = self.check_against_truth(df, wb)
        self.assertEqual(p[NAMA[1]][2], p[NAMA[2]][2])           # jam sama (data memang sama)
        self.assertNotEqual(p[NAMA[1]][3], p[NAMA[2]][3])         # tetapi jumlahnya dihitung sendiri-sendiri

    def test_kamera_tanpa_event(self):
        df = make_df({1: {8: (4, 1)}, 2: {9: (6, 2)}})
        _, wb = self.build(df, "kosong")
        t, p = self.check_against_truth(df, wb)
        self.assertEqual(p[NAMA[3]][2], rb.NA)
        self.assertEqual(p[NAMA[3]][4], rb.NA)
        self.assertEqual(wb["REKAP PER KAMERA"].cell(9, 11).value, rb.NA)

    def test_seri_dalam_satu_kamera_pilih_jam_paling_awal(self):
        df = make_df({1: {7: (10, 0), 9: (10, 0), 12: (3, 0), 14: (3, 0)}, 2: {10: (1, 0)}})
        _, wb = self.build(df, "seri")
        _, p = self.check_against_truth(df, wb)
        self.assertEqual(p[NAMA[1]][2], "07:00 – 07:59")
        self.assertEqual(p[NAMA[1]][4], "12:00 – 12:59")

    def test_jam_kosong_tidak_dianggap_sepi(self):
        df = make_df({1: {7: (9, 0), 12: (4, 0), 17: (2, 0)}, 2: {8: (1, 0)}})
        _, wb = self.build(df, "sepi")
        _, p = self.check_against_truth(df, wb)
        self.assertEqual(p[NAMA[1]][4], "17:00 – 17:59")          # bukan 00:00 (0 event); 17:00 punya 2 event => itu jam sepi sesuai definisi
        self.assertEqual(p[NAMA[1]][5], 2)

    def test_kasus_nyata_menyala_0854_mengikuti_definisi(self):
        """Semua kamera baru menyala 08:54. Sesuai definisi laporan (jam berevent terendah), jam 08 memang
        yang terendah; hasil TIDAK dipaksa berbeda, dan jam parsial ditandai di AUDIT KAMERA."""
        rows, i = [], 0
        for cid, per in {1: {8: 8, 9: 92, 10: 134, 11: 154, 12: 134, 13: 83},
                         2: {8: 8, 9: 38, 10: 68, 11: 63, 12: 94, 13: 38},
                         3: {8: 5, 9: 28, 10: 41, 11: 31, 12: 74, 13: 27}}.items():
            for jam, n in per.items():
                for k in range(n):
                    i += 1
                    mnt = 54 + k % 6 if jam == 8 else (k * 7) % 60
                    rows.append((i, datetime(2026, 10, 7, jam, mnt, 0), i, "MASUK", i, 0, cid, None))
        df = pd.DataFrame(rows, columns=COLS)
        _, wb = self.build(df, "parsial")
        _, p = self.check_against_truth(df, wb)
        self.assertEqual(p[NAMA[1]][2], "11:00 – 11:59")
        self.assertEqual(p[NAMA[2]][2], "12:00 – 12:59")
        self.assertEqual(p[NAMA[3]][2], "12:00 – 12:59")
        ws = wb["AUDIT KAMERA"]
        self.assertIn("08.00 (pertama", str(ws.cell(7, 29).value))

    def test_audit_kamera_status_dan_selisih_nol(self):
        df = make_df({1: {8: (5, 1), 9: (30, 15)}, 2: {9: (12, 3), 10: (50, 35)}})
        _, wb = self.build(df, "audit")
        ws = wb["AUDIT KAMERA"]
        t = truth(df)
        # tabel A: 24 jam per kamera = hitungan independen pandas
        for i, nm in enumerate(NAMA.values()):
            self.assertEqual(ws.cell(7 + i, 1).value, nm)
            self.assertEqual([ws.cell(7 + i, 4 + h).value for h in range(24)], [t[nm]["per_hour"].get(h, 0) for h in range(24)])
        hdr = next(r for r in range(1, 40) if ws.cell(r, 1).value == "Kamera" and ws.cell(r, 2).value == "Event Valid")
        for i, nm in enumerate(NAMA.values()):
            self.assertEqual(ws.cell(hdr + 1 + i, 12).value, 0, nm)
            self.assertEqual(ws.cell(hdr + 1 + i, 15).value, "SESUAI", nm)
            self.assertEqual(ws.cell(hdr + 1 + i, 2).value, t[nm]["event"], nm)

    def test_unit_compute_camera_stats_sintetis(self):
        from datetime import datetime as D
        mk = lambda n, c, h, st="MASUK": {"cam": c, "ts": D(2026, 1, 1, h, 0), "status": st}
        recs = [mk(0, "A", 9)] * 5 + [mk(0, "A", 12)] * 2 + [mk(0, "A", 14)] * 2 + [mk(0, "B", 10)] * 3 + [mk(0, "B", 11, "KELUAR")]
        recs.append({"cam": "A", "ts": None, "status": "MASUK"})
        st = rb.compute_camera_stats(recs, [{"name": "A", "cid": 1}, {"name": "B", "cid": 2}, {"name": "C", "cid": 3}])
        self.assertEqual(st["A"]["ramai"], (9, 5))
        self.assertEqual(st["A"]["sepi"], (12, 2))          # seri 12 & 14 -> paling awal
        self.assertEqual(st["A"]["valid"], 9)           # event tanpa timestamp tidak dihitung
        self.assertEqual(st["B"]["ramai"], (10, 3))
        self.assertEqual(st["B"]["sepi"], (11, 1))
        self.assertEqual((st["B"]["masuk"], st["B"]["keluar"]), (3, 1))
        self.assertIsNone(st["C"]["ramai"])

    def test_data_cacat_tidak_menyamakan_atau_merusak_kamera(self):
        extra = [(None, 5, "MASUK", 1, 1, 1, None),                               # tanpa timestamp
                 (datetime(2026, 10, 7, 11, 0, 0), 6, "MASUK", 1, 1, None, None),  # tanpa kamera
                 (datetime(2026, 10, 7, 11, 5, 0), 7, None, 1, 1, 2, None),        # tanpa status
                 (datetime(2026, 10, 7, 11, 6, 0), 8, "LAIN", 1, 1, 2, None)]      # status asing
        df = make_df({1: {8: (6, 2)}, 2: {9: (9, 3)}}, extra)
        _, wb = self.build(df, "cacat")
        self.check_against_truth(df, wb)
        ws = wb["VALIDASI"]
        txt = {ws.cell(r, 2).value: ws.cell(r, 3).value for r in range(7, 15)}
        self.assertEqual(txt["Timestamp kosong"], 1)
        self.assertEqual(txt["Kamera kosong"], 1)
        self.assertEqual(txt["Status kosong"], 1)
        self.assertEqual(txt["Status di luar MASUK/KELUAR"], 1)
        status = next(ws.cell(r, 3).value for r in range(15, 40) if ws.cell(r, 2).value == "STATUS KESELURUHAN")
        self.assertEqual(status, "PERLU VALIDASI DATA")           # selisih TIDAK disembunyikan

    def test_konsistensi_detail_rekap_dashboard(self):
        df = make_df({1: {8: (5, 1), 9: (30, 15)}, 2: {10: (50, 35)}})
        _, wb = self.build(df, "konsisten")
        ws = wb["VALIDASI"]
        status = next(ws.cell(r, 3).value for r in range(15, 40) if ws.cell(r, 2).value == "STATUS KESELURUHAN")
        self.assertEqual(status, "SEMUA ANGKA KONSISTEN")
        self.assertEqual(wb["DASHBOARD"]["B14"].value, 5 + 30 + 50)
        self.assertEqual(wb["DASHBOARD"]["D14"].value, 1 + 15 + 35)

    def test_grafik_menunjuk_kolom_kamera_yang_benar(self):
        df = make_df({1: {8: (5, 1)}, 2: {9: (3, 3)}})
        path, _ = self.build(df, "grafik")
        with zipfile.ZipFile(path) as z:
            xml = "".join(z.read(n).decode("utf8") for n in z.namelist() if n.startswith("xl/charts/chart"))
        refs = re.findall(r"<(?:c:)?f>([^<]*)</(?:c:)?f>", xml)
        ana = [r.replace("$", "").replace("'", "") for r in refs if "ANALISIS" in r]
        # grafik 'Event per Jam per Kamera': satu seri per kamera, judul seri = header kolom kamera,
        # data = kolom kamera yang SAMA (B, C, D) -> bukan satu kolom yang sama untuk semua kamera
        titles = [r for r in ana if re.fullmatch(r"ANALISIS!([A-Z]+)7", r)]
        data = [r for r in ana if re.fullmatch(r"ANALISIS!([A-Z]+)8:\1 ?31", r.replace("$", ""))]
        data = [r for r in ana if re.fullmatch(r"ANALISIS![A-Z]+8:[A-Z]+31", r)]
        data_cols = [re.match(r"ANALISIS!([A-Z]+)8:([A-Z]+)31", r).groups() for r in data if not r.startswith("ANALISIS!A8")]
        self.assertEqual([t.split("!")[1] for t in titles[:3]], ["B7", "C7", "D7"])
        self.assertEqual([c[0] for c in data_cols[:3]], ["B", "C", "D"])
        self.assertTrue(all(a == b for a, b in data_cols[:3]))

    def test_dashboard_filter_kamera_mengubah_hasil(self):
        df = make_df({1: {8: (5, 1), 9: (30, 15)}, 2: {9: (12, 3), 10: (50, 35)}})
        path = os.path.join(self.tmp, "filter.xlsx")
        rb.build_report_workbook(path, df, CAMS, date(2026, 10, 7), date(2026, 10, 7), snapshot_dir=None)
        wb = load_workbook(path)
        wb["DASHBOARD"]["C7"] = NAMA[2]
        wb.save(path)
        w = recalc(path)["DASHBOARD"]
        self.assertEqual([w[c].value for c in ("B18", "D18", "F18")], [62, 38, 100])

    def test_rekap_iv_filter_dan_diagram_mengikuti_kamera(self):
        df = make_df({1: {8: (5, 1), 9: (30, 15)}, 2: {9: (12, 3), 10: (50, 35)}, 3: {11: (4, 0)}})
        path = os.path.join(self.tmp, "vi.xlsx")
        rb.build_report_workbook(path, df, CAMS, date(2026, 10, 7), date(2026, 10, 7), snapshot_dir=None)
        t = truth(df)
        for pilih in [None, NAMA[1], NAMA[2], NAMA[3]]:
            wbk = load_workbook(path)
            if pilih:
                wbk["REKAP"]["K6"] = pilih
                wbk["DASHBOARD"]["C7"] = pilih
            wbk.save(path + ".t.xlsx")
            wb = recalc(path + ".t.xlsx")
            sub = df if pilih is None else df[df["camera_id"] == [k for k, v in NAMA.items() if v == pilih][0]]
            hrs = pd.to_datetime(sub["timestamp"]).dt.hour
            for h in range(24):
                row = [wb["REKAP"].cell(13 + h, c).value for c in range(10, 16)]
                self.assertEqual(row[2:], [int(((hrs == h) & (sub.status == "MASUK")).sum()), int(((hrs == h) & (sub.status == "KELUAR")).sum()),
                                           int((hrs == h).sum()), int((hrs == h).sum())], f"{pilih} jam {h}")
            self.assertEqual(wb["REKAP"]["S12"].value, "Diagram People Counting — " + (pilih or "Semua Kamera"))
            d = wb["DASHBOARD"]
            self.assertEqual(d["F18"].value, len(sub))                       # tabel dashboard
            self.assertEqual(d["K14"].value, "Diagram People Counting — " + (pilih or "Semua Kamera") + " (per Jam)")

    def test_seluruh_sheet_tetap_ada(self):
        df = make_df({1: {8: (2, 1)}})
        path, wb = self.build(df, "sheet")
        for nm in ["DASHBOARD", "DATA DETAIL", "REKAP", "REKAP PER KAMERA", "ANALISIS", "AUDIT KAMERA", NAMA[1], NAMA[2],
                   NAMA[3], "FILTER KAMERA", "VALIDASI", "REFERENSI"]:
            self.assertIn(nm, wb.sheetnames)
        # tidak ada sel error di seluruh workbook
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("#"):
                        self.fail(f"Error rumus {c.value} di {ws.title}!{c.coordinate}")


if __name__ == "__main__":
    unittest.main(verbosity=2)
